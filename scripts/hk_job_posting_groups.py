"""Create two reproducible, overlapping job packets from the 300-job HK sample.

The input export stays read only. Output is a new directory containing one XLSX,
two PDFs, and a manifest with the seed, allocations, page references, and hashes.
"""
import argparse
from datetime import datetime, timezone
from html import escape
from html.parser import HTMLParser
import importlib.metadata
import json
from pathlib import Path
import random
import re
import secrets
import unicodedata

import hk_job_postings_sample as sample

ROOT = Path(__file__).resolve().parents[1]
JOB_HEADING = re.compile(r'\bJob\s+ID:\s*(-?[0-9]+)\b')
HTML_TAG = re.compile(r'</?(?:p|br|div|ul|ol|li|strong|b|em|i|u|h[1-6]|span|a|'
                      r'table|blockquote|script|style)\b', re.IGNORECASE)


class DescriptionHTML(HTMLParser):
    """Keep text and basic formatting; never allow active content or attributes."""

    allowed = {'p', 'br', 'div', 'ul', 'ol', 'li', 'strong', 'b', 'em', 'i', 'u',
               'h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'span', 'a', 'blockquote',
               'table', 'thead', 'tbody', 'tfoot', 'tr', 'td', 'th', 'hr', 'pre'}
    blocked = {'script', 'style', 'iframe', 'object', 'embed', 'svg', 'head'}
    blocks = {'p', 'br', 'div', 'li', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6',
              'blockquote', 'tr', 'td', 'th', 'hr', 'pre'}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.html = []
        self.text = []
        self.hidden = []

    def handle_starttag(self, tag, attrs):
        if tag in self.blocked:
            self.hidden.append(tag)
        if self.hidden:
            return
        if tag in self.allowed:
            self.html.append(f'<{tag}>')
        if tag in self.blocks:
            self.text.append('\n')
        if tag == 'img':
            alt = dict(attrs).get('alt', '')
            self.html.append(escape(alt))
            self.text.append(alt)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag in self.blocked:
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if self.hidden:
            if tag == self.hidden[-1]:
                self.hidden.pop()
            return
        if tag in self.allowed and tag not in {'br', 'hr'}:
            self.html.append(f'</{tag}>')
        if tag in self.blocks:
            self.text.append('\n')

    def handle_data(self, data):
        if not self.hidden:
            self.html.append(escape(data))
            self.text.append(data)


def description_content(text):
    # These private-use characters are legacy list bullets in the source ads.
    text = text.replace('\uf097', '•').replace('\uf0b7', '•')
    if not HTML_TAG.search(text):
        return escape(text), text
    parser = DescriptionHTML()
    parser.feed(text)
    parser.close()
    html = ''.join(parser.html)
    # Scraped ads often insert BRs around every block. Keep normal HTML spacing
    # instead of printing several blank lines between each list item.
    html = re.sub(r'<(p|div)>(?:\s|<br>)*</\1>', '', html)
    blocks = r'p|div|ul|ol|li|h[1-6]|blockquote|table|thead|tbody|tfoot|tr'
    html = re.sub(r'(?:<br>\s*)+(?=</?(?:' + blocks + r')>)', '', html)
    html = re.sub(r'(</?(?:' + blocks + r')>)\s*(?:<br>\s*)+', r'\1', html)
    html = re.sub(r'(?:<br>\s*)+(?=</(?:strong|b|em|i|u|span)>)', '', html)
    html = re.sub(r'(?:<br>\s*){3,}', '<br><br>', html)
    return html, ''.join(parser.text)


def allocate(ids, seed):
    if len(ids) != 300 or len(set(ids)) != 300:
        raise ValueError('Expected exactly 300 distinct job IDs.')
    shuffled = sorted(ids)
    random.Random(seed).shuffle(shuffled)
    shared = shuffled[:100]
    return [shared + shuffled[100:200], shared + shuffled[200:300]]


def validate_allocation(groups):
    if len(groups) != 2 or any(len(g) != 200 or len(set(g)) != 200 for g in groups):
        raise ValueError('Each group must have 200 distinct jobs.')
    if (groups[0][:100] != groups[1][:100]
            or set(groups[0]) & set(groups[1]) != set(groups[0][:100])
            or len(set(groups[0]) | set(groups[1])) != 300):
        raise ValueError('Overlap, shared ordering, or union is incorrect.')


def packet_html(records):
    sections = []
    for record in records:
        body, _ = description_content(record['description'])
        company = record['posting'].get('company')
        metadata = ([company] if company else []) + [record['location_raw']]
        sections.append(
            '<section class="job"><header>'
            f'<div class="job-id">Job ID: {escape(record["job_id"])}</div>'
            f'<h1>{escape(record["title_raw"])}</h1>'
            + ''.join(f'<div class="metadata">{escape(s)}</div>' for s in metadata)
            + f'</header><div class="description">{body}</div></section>')
    return '''<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'">
<title>Hong Kong Job Postings</title><style>
@page { size: A4; }
* { box-sizing: border-box; }
body { margin: 0; color: #17212b; font: 11pt/1.45 Arial, 'Arial Unicode MS', 'PingFang HK',
       'Noto Sans CJK TC', sans-serif; overflow-wrap: anywhere; }
.job + .job { break-before: page; }
header { break-inside: avoid; border-bottom: 1pt solid #71808e;
         padding-bottom: 12pt; margin-bottom: 15pt; }
.job-id { font-size: 12pt; font-weight: bold; color: #203a51; }
h1 { font-size: 18pt; line-height: 1.2; margin: 8pt 0; }
.metadata { margin-top: 3pt; }
.description { white-space: pre-wrap; }
.description p, .description div { margin: 0 0 8pt; }
.description ul, .description ol { white-space: normal; padding-left: 20pt; }
.description li { white-space: pre-wrap; margin-bottom: 4pt; }
.description h1, .description h2, .description h3,
.description h4, .description h5, .description h6 {
  font-size: 12pt; margin: 10pt 0 6pt; break-after: avoid; }
p, li { orphans: 2; widows: 2; }
table { width: 100%; border-collapse: collapse; table-layout: fixed; }
td, th { padding: 3pt; vertical-align: top; }
pre { white-space: pre-wrap; font: inherit; }
</style></head><body>''' + ''.join(sections) + '</body></html>'


def render_packet(page, records, path):
    page.set_content(packet_html(records), wait_until='load')
    page.evaluate('document.fonts.ready')
    page.pdf(path=str(path), format='A4', print_background=True,
             margin={'top': '18mm', 'bottom': '20mm', 'left': '18mm', 'right': '18mm'},
             display_header_footer=True, header_template='<span></span>',
             footer_template='''<div style="width:100%;text-align:center;
               font:9px Arial;color:#56616c">Page <span class="pageNumber"></span>
               of <span class="totalPages"></span></div>''')


def normalized_words(text):
    """Compare all letters/numbers, tolerating PDF whitespace and glyph encoding."""
    return ''.join(c for c in unicodedata.normalize('NFKC', text) if c.isalnum())


def inspect_pdf(path, ids, records):
    from pypdf import PdfReader

    reader = PdfReader(path)
    starts, chunks = [], []
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text()
        if not text.strip():
            raise ValueError(f'Blank page in {path.name}: {number}')
        headings = JOB_HEADING.findall(text)
        if len(headings) > 1:
            raise ValueError(f'Multiple job headings on page {number}.')
        if headings:
            starts.append({'job_id': headings[0], 'page_number': number})
            chunks.append([])
        if not chunks:
            raise ValueError('PDF begins without a job heading.')
        # Strip the generated footer before testing complete description text.
        text = re.sub(r'Page\s+\d+\s+of\s+\d+\s*$', '', text)
        chunks[-1].append(text)
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        if abs(width - 595.28) > 2 or abs(height - 841.89) > 2:
            raise ValueError('PDF page is not A4.')
    if [entry['job_id'] for entry in starts] != ids:
        raise ValueError('PDF job order differs from the allocation.')
    for entry, pages in zip(starts, chunks):
        _, original = description_content(records[entry['job_id']]['description'])
        if normalized_words(original) not in normalized_words('\n'.join(pages)):
            raise ValueError(f'Incomplete PDF description: {entry["job_id"]}')
    return starts, len(reader.pages)


def write_workbook(path, groups):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)
    for number, entries in enumerate(groups, 1):
        sheet = workbook.create_sheet(f'Group {number}')
        sheet.append(['job_id', 'page_number'])
        for index, entry in enumerate(entries):
            sheet.append([entry['job_id'], entry['page_number']])
            sheet.cell(index + 2, 1).number_format = '@'
            if index < 100:
                for cell in sheet[index + 2]:
                    cell.fill = PatternFill('solid', fgColor='EAF2F8')
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='203A51')
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions['A'].width = 24
        sheet.column_dimensions['B'].width = 18
    workbook.save(path)


def verify_workbook(path, groups):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True)
    try:
        if workbook.sheetnames != ['Group 1', 'Group 2']:
            raise ValueError('Workbook sheets differ from expected groups.')
        for sheet, expected in zip(workbook, groups):
            rows = list(sheet.iter_rows())
            if [c.value for c in rows[0]] != ['job_id', 'page_number']:
                raise ValueError('Unexpected workbook columns.')
            if [[c.value for c in row] for row in rows[1:]] != [
                    [entry['job_id'], entry['page_number']] for entry in expected]:
                raise ValueError('Workbook page references differ from PDFs.')
            if any(row[0].data_type != 's' for row in rows[1:]):
                raise ValueError('Workbook job IDs must be text.')
    finally:
        workbook.close()


def generate(source, output, seed=None):
    from playwright.sync_api import sync_playwright

    source, output = Path(source).resolve(), Path(output).resolve()
    if output.is_relative_to(source):
        raise ValueError('Output must be outside the read-only source sample.')
    source_validation = sample.verify_export(source)
    source_hash = sample.digest(source / 'manifest.json')
    source_manifest = json.loads((source / 'manifest.json').read_text(encoding='utf-8'))
    records = {entry['job_id']: json.loads((source / entry['filename']).read_text(
        encoding='utf-8')) for entry in source_manifest['selected']}
    seed = secrets.randbits(64) if seed is None else seed
    groups = allocate(list(records), seed)
    validate_allocation(groups)
    output.mkdir(parents=True, exist_ok=False)
    manifest = {'status': 'incomplete', 'source': str(source),
                'source_manifest_sha256': source_hash, 'source_validation': source_validation,
                'seed': seed, 'algorithm': 'Shuffle sorted IDs with random.Random(seed); '
                'first 100 shared, next 100 Group 1 only, last 100 Group 2 only.',
                'page_number': '1-based starting page in the corresponding PDF',
                'shared_ids': groups[0][:100], 'groups': [],
                'software': {name: importlib.metadata.version(name)
                             for name in ('playwright', 'openpyxl', 'pypdf')},
                'generator_sha256': sample.digest(__file__)}
    sample.write_json(output / 'manifest.json', manifest)
    entries = []
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        manifest['browser_version'] = browser.version
        try:
            page = browser.new_page(java_script_enabled=False)
            page.route('**/*', lambda route: route.abort())
            for number, ids in enumerate(groups, 1):
                filename = f'group_{number}.pdf'
                print(f'Rendering {filename}: 200 jobs...', flush=True)
                render_packet(page, [records[job_id] for job_id in ids], output / filename)
                starts, count = inspect_pdf(output / filename, ids, records)
                entries.append(starts)
                manifest['groups'].append({'sheet': f'Group {number}', 'pdf': filename,
                                           'page_count': count, 'jobs': starts})
                sample.write_json(output / 'manifest.json', manifest)
                print(f'Validated {filename}: {count} pages, full descriptions.', flush=True)
        finally:
            browser.close()
    write_workbook(output / 'groups.xlsx', entries)
    verify_workbook(output / 'groups.xlsx', entries)
    sample.verify_export(source)
    if sample.digest(source / 'manifest.json') != source_hash:
        raise ValueError('Source manifest changed during generation.')
    manifest['artifacts'] = {name: {'sha256': sample.digest(output / name),
                                   'bytes': (output / name).stat().st_size}
                             for name in ('groups.xlsx', 'group_1.pdf', 'group_2.pdf')}
    manifest['validation'] = {'allocation': True, 'pdf_job_order': True,
                              'full_description_letters_and_numbers': True,
                              'workbook_page_references': True, 'source_unchanged': True}
    manifest['finished_at_utc'] = datetime.now(timezone.utc).isoformat()
    manifest['status'] = 'complete'
    sample.write_json(output / 'manifest.json', manifest)
    print(json.dumps({'output': str(output), 'seed': seed,
                      'pages': [g['page_count'] for g in manifest['groups']]}), flush=True)
    return manifest


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path, help='Verified 300-job JSON sample directory')
    parser.add_argument('--output', type=Path, help='New output directory; never overwritten')
    parser.add_argument('--seed', type=int, help='Reproduce a previous allocation')
    args = parser.parse_args(argv)
    stamp = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    generate(args.source, args.output or ROOT / 'data' / f'hk_posting_groups_{stamp}', args.seed)


if __name__ == '__main__':
    main()
